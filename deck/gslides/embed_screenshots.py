"""
Embed the app screenshots into the Google Slides deck built by build_deck.py.

The Slides API only inserts images from public URLs, and the screenshots must not be
shared publicly. So this exports the deck as PPTX, swaps each "[print: X]" frame for the
matching PNG in deck/img/, and uploads the result as a new, private Google Slides file
(speaker notes, master and layouts survive the round-trip).

Usage:  python embed_screenshots.py <native_presentation_id>
Prints the new presentation id. Requires python-pptx.
"""
from __future__ import annotations

import json
import sys
import urllib.request
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Pt

from build_deck import QUOTA, token

IMG_DIR = Path(__file__).resolve().parents[1] / "img"
IMAGES = {"Painel": "painel.png", "Fila": "fila.png", "Agente": "agente.png", "Radar": "radar.png"}
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
NAME = "NIPs da ANS no prazo · NIP → VPP Copilot (FE Bar · Lucas Brabo)"


def main():
    pres = sys.argv[1]
    tmp_in, tmp_out = Path("/tmp/deck_native.pptx"), Path("/tmp/deck_with_images.pptx")

    # 1) export (the Slides export endpoint has no 10 MB Drive-export cap)
    req = urllib.request.Request(f"https://docs.google.com/presentation/d/{pres}/export/pptx")
    req.add_header("Authorization", f"Bearer {token()}")
    tmp_in.write_bytes(urllib.request.urlopen(req).read())

    # 2) swap frames for pictures (frame ratio == screenshot ratio, 1.6)
    deck = Presentation(str(tmp_in))
    for slide in deck.slides:
        for sh in list(slide.shapes):
            if sh.has_text_frame and sh.text_frame.text.startswith("[print:"):
                name = sh.text_frame.text.strip()[8:-1].strip()
                pic = slide.shapes.add_picture(str(IMG_DIR / IMAGES[name]), sh.left, sh.top, width=sh.width, height=sh.height)
                pic.line.color.rgb = RGBColor(0xDE, 0xD7, 0xCC)
                pic.line.width = Pt(1)
                sh._element.getparent().remove(sh._element)
    deck.save(str(tmp_out))

    # 3) resumable upload, converting back to Google Slides (private by default)
    meta = json.dumps({"name": NAME, "mimeType": "application/vnd.google-apps.presentation"}).encode()
    start = urllib.request.Request("https://www.googleapis.com/upload/drive/v3/files?uploadType=resumable&fields=id",
                                   data=meta, method="POST")
    for k, v in {"Authorization": f"Bearer {token()}", "x-goog-user-project": QUOTA,
                 "Content-Type": "application/json; charset=UTF-8", "X-Upload-Content-Type": PPTX}.items():
        start.add_header(k, v)
    session = urllib.request.urlopen(start).headers["Location"]
    put = urllib.request.Request(session, data=tmp_out.read_bytes(), method="PUT")
    put.add_header("Authorization", f"Bearer {token()}")
    put.add_header("Content-Type", PPTX)
    new_id = json.loads(urllib.request.urlopen(put).read())["id"]
    print(new_id)


if __name__ == "__main__":
    main()
